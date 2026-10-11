"""File generation: PDF, PNG, XLSX, CSV and the plain-text form all carry the same fields."""

from __future__ import annotations

import csv
import io
import re
import shutil
import subprocess
from datetime import date

import pytest
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from invoice_creator.presets import get_preset
from invoice_creator.render import (
    DEMO_FOOTER,
    invoice_lines,
    invoice_rows,
    money,
    render,
    render_text,
    safe_filename,
)

TODAY = date(2026, 10, 11)
OWNER = ("Grand Avenue Trading LLC", "500 Grand Blvd, Suite 210, Kansas City, MO 64106")


def pdf_text(data: bytes) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)


def test_pdf_is_real_and_extractable():
    inv = get_preset("reyes", TODAY).invoice
    r = render(inv, "pdf", *OWNER)
    assert r is not None
    assert r.data.startswith(b"%PDF")
    assert r.filename == "RP-10517.pdf"
    assert r.content_type == "application/pdf"
    text = pdf_text(r.data)
    for needle in (
        "Reyes Packaging",
        "Invoice No: RP-10517",
        f"Invoice Date: {inv.invoice_date.isoformat()}",
        f"Due Date: {inv.due_date.isoformat()}",
        "PO Number: 7744",
        "Total Due: $6,400.00 USD",
        "Bank Name: BBVA Mexico",
        # CLABE is not a label the backend parser knows, so it prints on a canonical "Bank account" line
        "Bank account: 012180001234567891  (CLABE)",
        "SWIFT: BCMRMXMMPYM",
        "Email: ar@reyespackaging.mx",
        "RFC: RPA230214001",
        "Lucía Reyes",
        DEMO_FOOTER,
    ):
        assert needle in text, needle
    # the vendor name is the first line (what the backend's vendor guesser reads first) and is
    # repeated as an explicit "From:" line, the label its deterministic parser keys on
    assert text.strip().splitlines()[:2] == ["Reyes Packaging", "From: Reyes Packaging"]


def test_pdf_handles_many_line_items_across_pages():
    inv = get_preset("brightline-1", TODAY).invoice.model_copy(deep=True)
    inv.line_items = [inv.line_items[0].model_copy(update={"description": f"Item {i}"}) for i in range(80)]
    inv.amount = inv.line_total
    r = render(inv, "pdf", *OWNER)
    reader = PdfReader(io.BytesIO(r.data))
    assert len(reader.pages) >= 2
    assert "Total Due:" in pdf_text(r.data)


def test_png_is_a_valid_image_with_the_invoice_text_layout():
    inv = get_preset("nova", TODAY).invoice
    r = render(inv, "png", *OWNER)
    assert r.data.startswith(b"\x89PNG\r\n\x1a\n")
    assert r.filename == "NT-0002.png" and r.content_type == "image/png"
    img = Image.open(io.BytesIO(r.data))
    assert img.format == "PNG" and img.mode == "RGB"
    assert img.size[0] == 1600
    # tall enough for every line at the OCR-friendly 54 px pitch
    assert img.size[1] >= 220 + 54 * len(invoice_lines(inv, *OWNER))
    # mostly white page with dark text (not blank, not inverted)
    grey = img.convert("L")
    hist = grey.histogram()
    white = sum(hist[240:])
    dark = sum(hist[:80])
    assert white > 0.7 * grey.size[0] * grey.size[1]
    assert dark > 1000


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract binary not installed")
@pytest.mark.parametrize("key", ["prairie", "apex"])
def test_png_is_legible_to_tesseract(key, tmp_path):
    """What the backend's OCR path (pytesseract, --psm 6) must get back from our PNGs: the vendor as the
    first line (no junk from decoration), the labels it keys on, and digit-exact identifiers."""
    inv = get_preset(key, TODAY).invoice
    r = render(inv, "png", *OWNER)
    png = tmp_path / r.filename
    png.write_bytes(r.data)
    out = subprocess.run(
        ["tesseract", str(png), "stdout", "--psm", "6"], capture_output=True, text=True, check=True
    ).stdout
    lines = [ln.strip() for ln in out.splitlines() if ln.strip()]
    assert lines[0] == inv.vendor_name
    for needle in (
        f"Invoice No: {inv.invoice_number}",
        f"Total Due: {money(inv.amount, inv.currency, with_code=True)}",
        f"Account Number: {inv.bank.account_number}",
        "Payment Details",
    ):
        assert needle in lines, (needle, lines)


def test_xlsx_rows():
    inv = get_preset("midwest", TODAY).invoice
    r = render(inv, "xlsx", *OWNER)
    assert r.data.startswith(b"PK\x03\x04")
    assert r.content_type.endswith("spreadsheetml.sheet")
    wb = load_workbook(io.BytesIO(r.data), read_only=True)
    rows = [[("" if c is None else str(c)) for c in row] for row in wb.active.iter_rows(values_only=True)]
    as_pairs = {row[0]: row[1] for row in rows if len(row) >= 2}
    assert rows[0][0] == "Midwest Pallets"
    assert as_pairs["Invoice Number"] == "MP-5533"
    assert as_pairs["Total Due"] == "$4,950.00 USD"
    assert as_pairs["Account Number"] == "001234567890"
    assert as_pairs["Routing Number"] == "101000019"
    assert as_pairs["EIN"] == "47-1234567"
    assert ["Item", "Qty", "Unit Price", "Amount"] in [row[:4] for row in rows]
    assert ["48x40 GMA pallets", "550", "$9.00", "$4,950.00"] in [row[:4] for row in rows]
    assert rows[-1][0] == DEMO_FOOTER


def test_csv_rows_match_xlsx_rows():
    inv = get_preset("brightline-2", TODAY).invoice
    r = render(inv, "csv", *OWNER)
    assert r.content_type == "text/csv" and r.filename == "BL-3313.csv"
    rows = list(csv.reader(io.StringIO(r.data.decode("utf-8"))))
    assert rows == invoice_rows(inv, *OWNER)
    pairs = {row[0]: row[1] for row in rows if len(row) >= 2}
    assert pairs["Invoice Number"] == "BL-3313"
    assert pairs["Total Due"] == "$4,900.00 USD"
    assert pairs["Description"] == "Janitorial supplies, second half"


def test_plain_text_form_has_label_value_lines_and_no_remit_to_label():
    inv = get_preset("warsaw", TODAY).invoice
    text = render_text(inv, *OWNER)
    lines = text.splitlines()
    assert lines[0] == "Warsaw Steel Sp. z o.o."
    assert "Invoice No: WS-2026-0931" in lines
    assert "Total Due: $12,000.00 USD" in lines
    assert "IBAN: DE89 3704 0044 0532 0130 00" in lines
    assert "SWIFT: COBADEFFXXX" in lines
    assert "NIP: 5252345678" in lines
    assert any(ln.startswith("Notes: URGENT") for ln in lines)
    assert lines[-1] == DEMO_FOOTER
    # "Remit To" as a section header would be read by the backend as the vendor name; never emit it
    assert not any(re.match(r"(?i)^\W*remit\s*to", ln) for ln in lines)
    # every label/value pair is "Label: value" on one line
    for label in ("Invoice Date", "Due Date", "Bank Name", "Email", "Phone"):
        assert any(ln.startswith(f"{label}:") or f"{label}:" in ln for ln in lines), label


def test_body_format_has_no_file():
    assert render(get_preset("reyes", TODAY).invoice, "body") is None


def test_optional_fields_are_omitted_cleanly():
    inv = get_preset("nova", TODAY).invoice.model_copy(deep=True)
    inv.po_number = ""
    inv.remit_to.tax_id = ""
    inv.notes = ""
    inv.line_items = []
    text = render_text(inv)
    assert "PO Number" not in text and "EIN" not in text and "Notes:" not in text and "Bill To" not in text
    assert "Subtotal: $7,200.00" in text  # falls back to the amount when there are no lines


@pytest.mark.parametrize(
    ("amount", "currency", "expected"),
    [(6400, "USD", "$6,400.00"), (1234.5, "EUR", "EUR 1,234.50"), (0, "USD", "$0.00")],
)
def test_money_formatting(amount, currency, expected):
    assert money(amount, currency) == expected


def test_safe_filename():
    assert safe_filename("RP-10517", "pdf") == "RP-10517.pdf"
    assert safe_filename("../etc/passwd", "csv") == "etc-passwd.csv"
    assert safe_filename("Invoice #12 (final)", "xlsx") == "Invoice-12-final.xlsx"
    assert safe_filename("", "png") == "invoice.png"
